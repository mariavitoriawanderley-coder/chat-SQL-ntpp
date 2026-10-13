"""
exportacao.py
=============
Exportação dos resultados da pesquisa para .xlsx (Etapa 4).

Gera uma planilha com TODOS os campos retornados pela fonte. É um "dump" de
dados (sem fórmulas), então os valores são gravados como vieram — evitando
converter códigos, CPF/CNPJ e números de NE em número (o que os corromperia).
Aplica formatação leve: cabeçalho destacado, painel congelado, autofiltro,
fonte profissional e larguras de coluna.

ESCALA: a planilha é escrita em modo STREAMING (`write_only`), gravando linha a
linha em disco temporário em vez de montar todas as células na memória. É o que
permite exportar centenas de milhares de linhas (ex.: os itens de gasto de uma
pesquisa com dezenas de milhares de empenhos) sem estourar memória nem o tempo
da requisição. Como o Excel não aceita mais de 1.048.576 linhas por aba, o
excedente continua na aba seguinte ("Empenhos (2)", "Empenhos (3)", …).
"""

import io
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Regex de caracteres de controle proibidos pelo formato .xlsx. Se não estiver
# disponível na versão do openpyxl, usamos um fallback equivalente.
try:  # pragma: no cover
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
except Exception:  # pragma: no cover
    import re as _re

    ILLEGAL_CHARACTERS_RE = _re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

_FONTE = "Arial"
_MAX_LARGURA = 60
_MIN_LARGURA = 8
_MAX_CARACTERES = 32767      # limite de caracteres por célula no Excel
_MAX_LINHAS_ABA = 1048576    # limite de linhas por aba (inclui o cabeçalho)
_AMOSTRA_LARGURA = 200       # linhas usadas para estimar a largura das colunas


def valor_celula(v: Any) -> Any:
    """
    Prepara um valor para a célula:
    - números/bool/None passam direto;
    - strings (e demais tipos serializados) têm caracteres de controle ILEGAIS
      removidos (evita IllegalCharacterError do openpyxl) e são truncadas ao
      limite do Excel.
    """
    if v is None or isinstance(v, (int, float, bool)):
        return v
    s = v if isinstance(v, str) else str(v)
    s = ILLEGAL_CHARACTERS_RE.sub("", s)
    if len(s) > _MAX_CARACTERES:
        s = s[:_MAX_CARACTERES]
    return s


def _larguras(registros: List[Dict[str, Any]], colunas: List[str]) -> List[float]:
    """Larguras aproximadas, medidas numa amostra (não custa varrer tudo)."""
    amostra = registros[:_AMOSTRA_LARGURA]
    larguras = []
    for col in colunas:
        largura = len(str(col))
        for reg in amostra:
            largura = max(largura, len(str(reg.get(col, ""))))
        larguras.append(max(_MIN_LARGURA, min(_MAX_LARGURA, largura + 2)))
    return larguras


def _abrir_aba(wb: Workbook, titulo: str, colunas: List[str], larguras: List[float]):
    """Cria uma aba já com cabeçalho formatado, congelamento e larguras."""
    ws = wb.create_sheet(title=titulo)
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    # Em modo streaming a "sheet view" é gravada junto da primeira linha: o
    # congelamento precisa estar definido ANTES do primeiro append.
    ws.freeze_panes = "A2"

    negrito = Font(name=_FONTE, bold=True, color="FFFFFF")
    preenchimento = PatternFill("solid", start_color="14508C")
    cabecalho = []
    for nome in colunas:
        c = WriteOnlyCell(ws, value=nome)
        c.font = negrito
        c.fill = preenchimento
        c.alignment = Alignment(vertical="center")
        cabecalho.append(c)
    ws.append(cabecalho)
    return ws


def _titulo_aba(base: str, indice: int) -> str:
    """'Empenhos', 'Empenhos (2)', … respeitando o limite de 31 caracteres."""
    if indice == 1:
        return base[:31]
    sufixo = f" ({indice})"
    return f"{base[: 31 - len(sufixo)]}{sufixo}"


def escrever_workbook(
    wb: Workbook,
    registros: Iterable[Dict[str, Any]],
    colunas: List[str],
    titulo: str = "Empenhos",
) -> Workbook:
    """
    Escreve os registros no workbook (em modo streaming), quebrando em novas
    abas ao atingir o limite de linhas do Excel. Devolve o próprio workbook.
    """
    lista = registros if isinstance(registros, list) else list(registros)
    larguras = _larguras(lista, colunas)

    aba = 1
    ws = _abrir_aba(wb, _titulo_aba(titulo, aba), colunas, larguras)
    escritas = 0          # linhas de dados na aba atual
    total = 0
    for reg in lista:
        if escritas >= _MAX_LINHAS_ABA - 1:      # -1 por causa do cabeçalho
            ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{escritas + 1}"
            aba += 1
            ws = _abrir_aba(wb, _titulo_aba(titulo, aba), colunas, larguras)
            escritas = 0
        ws.append([valor_celula(reg.get(col)) for col in colunas])
        escritas += 1
        total += 1

    if colunas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{escritas + 1}"
    return wb


def construir_workbook(
    registros: List[Dict[str, Any]],
    colunas: Optional[List[str]] = None,
    titulo: str = "Empenhos",
) -> Workbook:
    """Workbook (write_only) com uma aba de dados — pronto para salvar."""
    if colunas is None:
        colunas = list(registros[0].keys()) if registros else []

    wb = Workbook(write_only=True)
    # Em modo streaming não dá para estilizar célula a célula sem custo. As
    # células de dados saem sem estilo próprio, ou seja, usam a fonte de índice
    # 0 da pasta: trocá-la aplica a fonte padrão a todas de uma vez.
    try:  # pragma: no cover - detalhe interno do openpyxl
        wb._fonts[0] = Font(name=_FONTE)
    except Exception:  # pragma: no cover - formatação é acessório
        pass

    escrever_workbook(wb, registros, colunas, titulo=titulo)
    return wb


def exportar_xlsx(
    registros: List[Dict[str, Any]],
    colunas: Optional[List[str]] = None,
    caminho: Optional[str] = None,
    titulo: str = "Empenhos",
) -> Any:
    """
    Gera a planilha. Se `caminho` for informado, salva e devolve o caminho;
    caso contrário, devolve os bytes do arquivo (útil para download via web).
    """
    wb = construir_workbook(registros, colunas, titulo=titulo)
    if caminho:
        wb.save(caminho)
        return caminho
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
